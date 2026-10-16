#!/usr/bin/env python3
"""Build the Excel workbook from the data package (run after tools/build.py).  Requires: pip install openpyxl"""
import csv, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
META = json.load(open(os.path.join(ROOT, 'tools', 'package_meta.json'), encoding='utf-8'))
OUT = os.path.join(ROOT, sys.argv[1] if len(sys.argv) > 1 else f'India_Govt_Blockchain_Database_v{META["version"].rsplit(".", 1)[0]}.xlsx')
def read(p): return list(csv.DictReader(open(os.path.join(ROOT, p), encoding='utf-8')))

F = 'Arial'; HDR = PatternFill('solid', fgColor='1F3864'); HF = Font(name=F, bold=True, color='FFFFFF', size=10); BF = Font(name=F, size=10)
thin = Side(style='thin', color='BFBFBF'); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
STAGE_FILL = {'Implemented': 'C6EFCE', 'In progress': 'DDEBF7', 'Proposed / Planned': 'FCE4D6', 'Pilot completed': 'FFE5CC', 'Unclear (no recent update)': 'EDEDED', 'Discontinued': 'F8CBAD'}
VOC = {}
for v in read('data/vocabularies.csv'): VOC.setdefault(v['vocabulary'], []).append(v['code'])
STAGES = VOC['stage']

def table_sheet(wb, name, rows, widths=None, freeze='B2', stage_col=None, wrap=()):
    ws = wb.create_sheet(name)
    hdr = list(rows[0].keys()) if rows else []
    ws.append(hdr)
    for c in ws[1]: c.font = HF; c.fill = HDR; c.border = BD; c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append([(int(v) if v.lstrip('-').isdigit() and len(v) < 15 and k not in ('lgd_code', 'project_id') else v) if isinstance(v, str) else v for k, v in r.items()])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BF; c.border = BD
            c.alignment = Alignment(wrap_text=hdr[c.column - 1] in wrap, vertical='top')
        if stage_col and row[hdr.index(stage_col)].value in STAGE_FILL:
            row[hdr.index(stage_col)].fill = PatternFill('solid', fgColor=STAGE_FILL[row[hdr.index(stage_col)].value])
    for i, h in enumerate(hdr, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, min(45, max(10, len(h) + 2)))
    ws.freeze_panes = freeze; ws.auto_filter.ref = ws.dimensions
    return ws

wb = Workbook(); ws = wb.active; ws.title = 'README'
P = read('data/all_projects.csv')
lines = [(META['title'], 14, True), (f'Version {META["version"]} · created {META["created"][:10]} · licence CC BY 4.0 · authors: Santanu Mondal (ORCID 0009-0001-6243-1709) and T. Chithralekha, Pondicherry University', 10, False), ('DOI (all versions, always the latest): https://doi.org/10.5281/zenodo.23136332 · this version 1.1.0: https://doi.org/10.5281/zenodo.23415599', 10, False), ('', 10, False),
 (META['description'], 10, False), ('', 10, False), ('Sheets', 11, True),
 ('All projects – THE COMBINED TABLE: every project (Union, State, UT, district, municipal, PSU, regulator, court ...) in one row with level of government, body type, region, who built it, stage dates, NIC usage, challenges and evidence', 10, False),
 ('Summary – live counts by stage, tier, use-case and evidence score (formulas on the Projects sheet)', 10, False),
 ('Stage events – dated lifecycle events with sources (announced, pilot, live, scaled, latest update)', 10, False),
 ('Challenges – source-reported challenges, TOE-coded (see Vocabularies for TOE/DOI mapping)', 10, False),
 ('Sources / Project sources – all sources with type and evidence score, and which project each supports', 10, False),
 ('NIC snapshots – time series of the NIC blockchain platform by platform, application and department', 10, False),
 ('Jurisdictions – all 36 States/UTs with ISO 3166-2 and LGD codes, search results and live project counts', 10, False),
 ('Bodies, Vocabularies – implementing bodies; controlled lists and mappings (COFOG, TOE, DOI, stage, region, body type, platform family)', 10, False),
 ('Reliability, Your Coding Sheet – consistency check (kappa) and a blind sheet for a human second coder (kappa calculates automatically)', 10, False),
 ('RTI drafts, Policies, Literature, Context stats, Excluded, Changelog – supporting tables', 10, False), ('', 10, False),
 ('The CSV files and datapackage.json are the master copy. Edit them, run tools/build.py (validation) and tools/make_workbook.py to rebuild this workbook.', 10, False),
 ('"Unclear" means no public update was found, not that a project failed.', 10, False)]
for i, (t, sz, b) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = Font(name=F, size=sz, bold=b); c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 130
for i in (4,): ws.row_dimensions[i].height = 44

W = {'project_name': 40, 'description': 50, 'status_text': 34, 'scale_text': 30, 'verify_note': 36, 'implementing_body': 32, 'platform_reported': 24,
     'tech_partner': 20, 'body_category': 26, 'implementation_model': 30, 'parent_ministry_or_state': 28, 'primary_source_url': 40, 'evidence_label': 34}
pw = table_sheet(wb, 'All projects', P, W, freeze='C2', stage_col='stage', wrap=('project_name', 'description', 'status_text', 'verify_note', 'scale_text'))
PH = list(P[0].keys()); col = lambda name: get_column_letter(PH.index(name) + 1); n = len(P) + 1
rng = lambda name: f"'All projects'!${col(name)}$2:${col(name)}${n}"
A = f'{rng("record_status")},"active"'

# Summary with formulas
sm = wb.create_sheet('Summary', 1)
sm['A1'] = "Summary (live formulas on the 'All projects' sheet; active records only)"; sm['A1'].font = Font(name=F, bold=True, size=13)
r = 3
def header(cells):
    global r
    for j, h in enumerate(cells, 1):
        c = sm.cell(row=r, column=j, value=h); c.font = HF; c.fill = HDR; c.border = BD
    r += 1
tiers = VOC['government_level']
T = get_column_letter(len(tiers) + 2)          # Total column
header(['Stage'] + tiers + ['Total', '% of all'])
s0 = r
for st in STAGES:
    sm.cell(row=r, column=1, value=st)
    for j, t in enumerate(tiers, 2): sm.cell(row=r, column=j, value=f'=COUNTIFS({rng("stage")},$A{r},{rng("government_level")},{get_column_letter(j)}$3,'+A+')')
    sm.cell(row=r, column=len(tiers) + 2, value=f'=SUM(B{r}:{get_column_letter(len(tiers) + 1)}{r})')
    sm.cell(row=r, column=len(tiers) + 3, value=f'=IF(${T}${s0 + len(STAGES)}=0,0,{T}{r}/${T}${s0 + len(STAGES)})').number_format = '0.0%'
    r += 1
sm.cell(row=r, column=1, value='Total').font = Font(name=F, bold=True, size=10)
for j in range(2, len(tiers) + 3): sm.cell(row=r, column=j, value=f'=SUM({get_column_letter(j)}{s0}:{get_column_letter(j)}{r - 1})')
tot = f'${T}${r}'; r += 2
impl, pil, unc = s0, s0 + STAGES.index('Pilot completed'), s0 + STAGES.index('Unclear (no recent update)')
sm.cell(row=r, column=1, value='Pilot-to-production rate = Implemented ÷ (Implemented + Pilot completed + Unclear)').font = Font(name=F, bold=True, size=10)
c = sm.cell(row=r, column=2, value=f'=IFERROR({T}{impl}/({T}{impl}+{T}{pil}+{T}{unc}),0)'); c.number_format = '0.0%'; r += 1
sm.cell(row=r, column=1, value='Share backed by an official source (evidence score 4–5)').font = Font(name=F, bold=True, size=10)
c = sm.cell(row=r, column=2, value=f'=COUNTIFS({rng("evidence_score")},">=4",'+A+f')/{tot}'); c.number_format = '0.0%'; r += 1
sm.cell(row=r, column=1, value='Mean evidence score').font = Font(name=F, bold=True, size=10)
c = sm.cell(row=r, column=2, value=f'=AVERAGEIFS({rng("evidence_score")},'+A+')'); c.number_format = '0.00'; r += 2
for title, colname, voc in [('Type of public body', 'body_category', 'body_category'), ('Built by', 'implementation_model', 'implementation_model'),
                            ('Region', 'region', 'region'), ('Use-case', 'use_case', 'use_case')]:
    header([title, 'Implemented', 'All', 'Implemented %'])
    for v in VOC[voc]:
        sm.cell(row=r, column=1, value=v)
        sm.cell(row=r, column=2, value=f'=COUNTIFS({rng(colname)},$A{r},{rng("stage")},"Implemented",'+A+')')
        sm.cell(row=r, column=3, value=f'=COUNTIFS({rng(colname)},$A{r},'+A+')')
        sm.cell(row=r, column=4, value=f'=IF(C{r}=0,"",B{r}/C{r})').number_format = '0%'; r += 1
    r += 1
header(['Evidence score', 'Projects'])
for k in range(5, 0, -1):
    sm.cell(row=r, column=1, value=k); sm.cell(row=r, column=2, value=f'=COUNTIFS({rng("evidence_score")},$A{r},'+A+')'); r += 1
for row in sm.iter_rows(min_row=3):
    for c in row:
        if c.value is not None and c.fill != HDR and not c.font.bold: c.font = BF
        if c.value is not None: c.border = BD
sm.column_dimensions['A'].width = 44
for L in 'BCDEFG': sm.column_dimensions[L].width = 22

table_sheet(wb, 'Stage events', read('data/stage_events.csv'), {'evidence': 70, 'recorded_by': 26}, wrap=('evidence',))
table_sheet(wb, 'Challenges', read('data/challenges.csv'), {'description': 50, 'evidence': 50}, wrap=('description', 'evidence'))
table_sheet(wb, 'Sources', read('data/sources.csv'), {'title': 60, 'author': 30, 'publisher': 30, 'url': 50}, wrap=('title',))
table_sheet(wb, 'Project sources', read('data/project_sources.csv'))
table_sheet(wb, 'NIC snapshots', read('data/nic_snapshots.csv'), {'unit': 50, 'unit_as_published': 50, 'note': 40}, freeze='A2')
J = read('data/jurisdictions.csv')
jw = table_sheet(wb, 'Jurisdictions', J, {'name': 30, 'search_result': 26, 'search_note': 60}, freeze='C2', wrap=('search_note',))
jw.cell(row=1, column=9, value='projects (live formula)'); jw.cell(row=1, column=10, value='implemented (live formula)')
for c in (jw.cell(row=1, column=9), jw.cell(row=1, column=10)): c.font = HF; c.fill = HDR; c.border = BD
for i in range(2, len(J) + 2):
    jw.cell(row=i, column=9, value=f'=COUNTIFS({rng("jurisdiction_id")},$A{i},'+A+')').border = BD
    jw.cell(row=i, column=10, value=f'=COUNTIFS({rng("jurisdiction_id")},$A{i},{rng("stage")},"Implemented",'+A+')').border = BD
jw.column_dimensions['I'].width = 14; jw.column_dimensions['J'].width = 14
table_sheet(wb, 'Bodies', read('data/bodies.csv'), {'name': 50, 'body_type': 26, 'parent': 34})
table_sheet(wb, 'Vocabularies', read('data/vocabularies.csv'), {'label': 50, 'definition': 70, 'maps_to': 34}, freeze='A2', wrap=('definition',))
table_sheet(wb, 'Reliability', read('data/reliability.csv'), freeze='A2')

# ---------- blind coding sheet (human second coder) ----------
SID = ['C01', 'C02', 'C11', 'C13', 'C14', 'C15', 'C16', 'S07', 'S09', 'S11', 'S17', 'S21', 'S31', 'S34', 'S37', 'S38', 'S43', 'S45', 'S46', 'S51',
       'S52', 'S54', 'S56', 'S57', 'S59', 'S60', 'S63', 'S74', 'S77', 'O01', 'O02', 'O03', 'O07', 'O13', 'O18']   # random 30% sample, seed 2026 (v1.0.0)
PM = {p['project_id']: p for p in P}
lists = [STAGES, VOC['use_case'], VOC['chain_type'], VOC['service_type'], VOC['scale_tier']]
CODEF = ['stage', 'use_case', 'chain_type', 'service_type', 'scale_tier']
ak = wb.create_sheet('_AnswerKey')
ak.append(['project_id'] + CODEF + [''] + ['list_' + c for c in CODEF])
for k, sid in enumerate(SID):
    rr = k + 2; ak.cell(row=rr, column=1, value=sid)
    for j, f in enumerate(CODEF):
        ak.cell(row=rr, column=2 + j, value=f'=INDEX({rng(f)},MATCH($A{rr},{rng("project_id")},0))')
for j, lst in enumerate(lists):
    for k, v in enumerate(lst): ak.cell(row=k + 2, column=8 + j, value=v)
ak.sheet_state = 'hidden'
yc = wb.create_sheet('Your Coding Sheet')
yc['A1'] = 'Blind coding sheet: be the human second coder'; yc['A1'].font = Font(name=F, bold=True, size=13)
yc['A2'] = ('Read each row and choose a value in the 5 yellow columns from the drop-downs, using the Vocabularies sheet. Do not look at the Projects sheet while coding. '
            'Agreement and Cohen\'s kappa against the dataset codes appear at the bottom automatically (about 40 minutes).')
yc['A2'].font = Font(name=F, italic=True, size=10); yc['A2'].alignment = Alignment(wrap_text=True); yc.merge_cells('A2:L2'); yc.row_dimensions[2].height = 32
H2 = ['ID', 'Govt body', 'Project', 'What it does', 'Blockchain platform (reported)', 'Current status', 'Scale (reported)', 'Stage', 'Use-case type', 'Chain type', 'Service type', 'Scale tier']
for j, h in enumerate(H2, 1):
    c = yc.cell(row=4, column=j, value=h); c.font = HF; c.fill = HDR if j < 8 else PatternFill('solid', fgColor='BF8F00'); c.border = BD; c.alignment = Alignment(wrap_text=True)
YF = PatternFill('solid', fgColor='FFF2CC')
for k, sid in enumerate(SID):
    rr = k + 5; p = PM[sid]
    for j, v in enumerate([sid, p['implementing_body'], p['project_name'], p['description'], p['platform_reported'], p['status_text'], p['scale_text']], 1):
        c = yc.cell(row=rr, column=j, value=v); c.font = BF; c.border = BD; c.alignment = Alignment(wrap_text=True, vertical='top')
    for j in range(8, 13): c = yc.cell(row=rr, column=j); c.fill = YF; c.border = BD
last = 4 + len(SID)
for j, (L, lst) in enumerate(zip('HIJKL', lists)):
    dv = DataValidation(type='list', formula1=f"'_AnswerKey'!${'HIJKL'[j]}$2:${'HIJKL'[j]}${len(lst) + 1}", allow_blank=True)
    yc.add_data_validation(dv); dv.add(f'{L}5:{L}{last}')
rr = last + 2
yc.cell(row=rr, column=1, value='Results (automatic)').font = Font(name=F, bold=True, size=11); rr += 1
for j, h in enumerate(['Variable', 'Rows coded', 'Agreement', "Cohen's kappa"], 1):
    c = yc.cell(row=rr, column=j, value=h); c.font = HF; c.fill = HDR; c.border = BD
rr += 1
for L, KL, cl, nm in zip('HIJKL', 'BCDEF', 'HIJKL', ['Stage', 'Use-case type', 'Chain type', 'Service type', 'Scale tier']):
    U = f'{L}$5:{L}${last}'; K = f"'_AnswerKey'!{KL}$2:{KL}${len(SID) + 1}"; C = f"'_AnswerKey'!{cl}$2:{cl}$20"
    yc.cell(row=rr, column=1, value=nm)
    yc.cell(row=rr, column=2, value=f'=COUNTA({U})')
    yc.cell(row=rr, column=3, value=f'=IF(B{rr}=0,"",SUMPRODUCT(({U}<>"")*({U}={K}))/B{rr})')
    pe = f'SUMPRODUCT(COUNTIFS({U},{C})*COUNTIFS({K},{C},{U},"<>"))/B{rr}^2'
    yc.cell(row=rr, column=4, value=f'=IF(B{rr}=0,"",IFERROR((C{rr}-{pe})/(1-{pe}),1))')
    yc.cell(row=rr, column=3).number_format = '0%'; yc.cell(row=rr, column=4).number_format = '0.00'
    for c in yc[rr][:4]: c.font = BF; c.border = BD
    rr += 1
for L, w in zip('ABCDEFGHIJKL', [7, 26, 34, 40, 20, 30, 24, 16, 22, 20, 10, 18]): yc.column_dimensions[L].width = w
yc.freeze_panes = 'C5'

for name, f, wd in [('RTI drafts', 'rti_drafts', {'draft_rti_application_copy_fill_send': 90}), ('Policies', 'policies', None), ('Literature', 'literature', {'title': 60}),
                    ('Context stats', 'context_stats', None), ('Excluded', 'excluded_items', None), ('Changelog', 'changelog', {'description': 90})]:
    rows = read(f'data/{f}.csv')
    table_sheet(wb, name, rows, wd, freeze='A2', wrap=tuple(rows[0].keys()) if name == 'RTI drafts' else ('description',))
wb.save(OUT); print('workbook:', OUT)
